"""Export Excel : le même lot d'offres, mais triable et annotable à la main."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .journal import obtenir
from .modeles import Offre

_log = obtenir("excel")

COLONNES = [
    ("Score", 8), ("Titre", 46), ("Entreprise", 26), ("Lieu", 22),
    ("Contrat", 16), ("Source", 14), ("Publiée le", 13),
    ("Pourquoi elle remonte", 46), ("Lien", 58),
]

ENTETE = PatternFill("solid", fgColor="1F3A5F")
ACCENT = PatternFill("solid", fgColor="FFF4D6")


def exporter(offres: list[Offre], dossier: Path | str = "exports") -> Path | None:
    """Écrit un classeur horodaté. Rend son chemin, ou None si rien à écrire."""
    if not offres:
        return None

    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = dossier / f"offres_{datetime.now():%Y-%m-%d_%Hh%M}.xlsx"

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Offres"

    for colonne, (intitule, largeur) in enumerate(COLONNES, start=1):
        cellule = feuille.cell(row=1, column=colonne, value=intitule)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = ENTETE
        cellule.alignment = Alignment(vertical="center")
        feuille.column_dimensions[get_column_letter(colonne)].width = largeur

    for ligne, offre in enumerate(offres, start=2):
        valeurs = (
            offre.score,
            offre.titre,
            offre.entreprise,
            offre.lieu,
            offre.contrat,
            offre.source,
            offre.date_publication.isoformat() if offre.date_publication else "",
            " ; ".join(offre.motifs),
            offre.url,
        )
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
            cellule.alignment = Alignment(vertical="top", wrap_text=colonne in (2, 8))

        # Le lien reste cliquable depuis Excel : c'est tout l'intérêt.
        if offre.url:
            lien = feuille.cell(row=ligne, column=len(COLONNES))
            lien.hyperlink = offre.url
            lien.font = Font(color="0563C1", underline="single")

        # Les meilleures offres se repèrent sans lire la colonne Score.
        if offre.score >= 20:
            for colonne in range(1, len(COLONNES) + 1):
                feuille.cell(row=ligne, column=colonne).fill = ACCENT

    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{len(offres) + 1}"
    classeur.save(chemin)

    _log.info("export Excel : %s (%d offres)", chemin, len(offres))
    return chemin
